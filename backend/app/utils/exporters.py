"""Server-side Export Generators for CSV, Excel (.xlsx), and PDF formats."""

import csv
import io
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def export_to_csv(headers: List[str], rows: List[List[Any]]) -> bytes:
    """Generate CSV byte stream with UTF-8 BOM for seamless Microsoft Excel opening."""
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    # UTF-8 BOM (\xef\xbb\xbf) ensures Microsoft Excel and Windows immediately recognize and open CSV files
    return b"\xef\xbb\xbf" + output.getvalue().encode("utf-8")


def export_to_excel(sheet_name: str, headers: List[str], rows: List[List[Any]]) -> bytes:
    """Generate professional Excel workbook as bytes with styled header and auto column widths."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel limits sheet title to 31 chars

    # Write headers
    ws.append(headers)

    # Style header row
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy blue
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Write data rows
    data_font = Font(name="Arial", size=10)
    for row in rows:
        ws.append(row)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.font = data_font
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")

    # Adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_to_pdf(
    title: str,
    headers: List[str],
    rows: List[List[Any]],
    summary: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Generate professional PDF report using ReportLab with table layout."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30,
    )
    elements = []
    styles = getSampleStyleSheet()

    # Title
    title_style = ParagraphStyle(
        name="ReportTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1E3A8A"),
        spaceAfter=10,
    )
    elements.append(Paragraph(f"VIT Mithra Attendance Portal - {title}", title_style))

    # Optional summary block
    if summary:
        summary_text = " | ".join([f"<b>{k}:</b> {v}" for k, v in summary.items()])
        summary_style = ParagraphStyle(
            name="SummaryStyle",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#374151"),
            spaceAfter=12,
        )
        elements.append(Paragraph(summary_text, summary_style))

    elements.append(Spacer(1, 10))

    # Table content
    table_data = [headers] + [[str(c) for c in row] for row in rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 10),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("TOPPADDING", (0, 0), (-1, 0), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 1), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    elements.append(table)

    doc.build(elements)
    return buffer.getvalue()
